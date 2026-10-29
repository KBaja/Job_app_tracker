"""
Job application tracker

Folder layout:
    Job_market/                 <- your tracker .xlsx and application folders
        Script/
            job_tracker.pyw     <- this file
            settings.json       <- created when you save settings
            assets/
                icon_setting.png
                celebrate.wav, interview.wav, offer.wav

Requires: openpyxl, customtkinter, pillow
"""
import os
import re
import sys
import json
import random
import shutil
import calendar
import subprocess
import datetime as dt
import tkinter as tk
from copy import copy, deepcopy
from tkinter import ttk, messagebox, filedialog

import customtkinter as ctk
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font

# =====================================================================
# Settings
# =====================================================================
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SETTINGS_PATH = os.path.join(SCRIPT_DIR, "settings.json")
ASSETS = os.path.join(SCRIPT_DIR, "assets")
DATE_FMT = "%d.%m.%Y"

PATTERN_PRESETS = [
    "{dd}_{mm}_{company}",
    "{yyyy}_{mm}_{dd}_{company}",
    "{yyyy}_{mm}_{dd}_{company}_{position}",
]

DEFAULT_SETTINGS = {
    "job_dir": "",              # empty = the folder above Script
    "excel_file": "",           # empty = the only .xlsx in job_dir
    "monthly_goal": 10,
    "appearance": "System",     # System, Light, Dark
    "defaults": {
        "workload": "100%",
        "method": "Online",
        "source": "LinkedIn",
        "location": "Zurich",
    },
    "status_options": ["Applied", "Interview", "Offer", "Rejected", "No response", "Withdrawn"],
    "rav_options": ["Pending", "Interview", "Rejected", "Hired"],
    "folder_pattern": PATTERN_PRESETS[0],
    "templates": [],
    "followup_days": 14,
    "sound": True,
    "sound_file": "assets/celebrate.wav",
    "sound_interview": "assets/interview.wav",
    "sound_offer": "assets/offer.wav",
    "click_sounds": True,
}


def load_settings():
    s = deepcopy(DEFAULT_SETTINGS)
    if os.path.exists(SETTINGS_PATH):
        try:
            with open(SETTINGS_PATH, encoding="utf-8") as f:
                user = json.load(f)
            for k, v in user.items():
                if isinstance(v, dict) and isinstance(s.get(k), dict):
                    s[k].update(v)
                else:
                    s[k] = v
        except Exception:
            pass
    return s


def save_settings():
    with open(SETTINGS_PATH, "w", encoding="utf-8") as f:
        json.dump(S, f, indent=2, ensure_ascii=False)


S = load_settings()


def script_rel(p):
    return p if os.path.isabs(p) else os.path.join(SCRIPT_DIR, p)


def job_dir():
    return S["job_dir"] or os.path.dirname(SCRIPT_DIR)


# =====================================================================
# Excel layer
# =====================================================================
# Semantic fields, matched against your header names by prefix
FIELDS = {
    "id": ("id",),
    "date": ("date applied", "date"),
    "company": ("company",),
    "address": ("address",),
    "contact": ("contact person",),
    "email": ("contact email", "contact e", "contact phone"),
    "position": ("position",),
    "workload": ("workload", "pensum"),
    "method": ("application method", "method"),
    "source": ("source",),
    "link": ("job ad link", "job link", "link"),
    "location": ("location",),
    "status": ("status",),
    "last": ("last update",),
    "result": ("result",),
    "folder": ("folder",),
    "notes": ("notes",),
}
DATE_FIELDS = ("date", "last")


def find_excel():
    if S["excel_file"]:
        p = S["excel_file"]
        return p if os.path.isabs(p) else os.path.join(job_dir(), p)
    d = job_dir()
    files = [f for f in os.listdir(d) if f.lower().endswith(".xlsx") and not f.startswith("~$")]
    if len(files) != 1:
        raise RuntimeError(f"Expected exactly one .xlsx in\n{d}\nfound {len(files)}. "
                           "Choose the tracker file in Settings.")
    return os.path.join(d, files[0])


def ensure_closed(path):
    if not os.path.exists(path):
        return
    try:
        with open(path, "a"):
            pass
    except PermissionError:
        raise RuntimeError(f"{os.path.basename(path)} is open in Excel. Close it and try again.")


def sheet(wb):
    return wb.worksheets[0]


def field_of(header):
    h = header.lower().strip()
    for f, prefixes in FIELDS.items():
        if any(h.startswith(p) for p in prefixes):
            return f
    return None


def find_field(headers, field):
    for p in FIELDS[field]:
        for col, h in headers:
            if h.lower().startswith(p):
                return col
    return None


def as_text(v):
    if v is None:
        return ""
    if isinstance(v, (dt.datetime, dt.date)):
        return v.strftime(DATE_FMT)
    return str(v)


def parse_date(s):
    try:
        return dt.datetime.strptime(s.strip(), DATE_FMT)
    except (ValueError, AttributeError):
        return None


def clean(name, maxlen=40):
    name = re.sub(r'[\\/:*?"<>|]', "", str(name)).strip()
    return re.sub(r"\s+", "_", name)[:maxlen]


def make_folder_name(date_str, company, position):
    d = parse_date(date_str) or dt.datetime.today()
    kw = dict(yyyy=f"{d:%Y}", yy=f"{d:%y}", mm=f"{d:%m}", dd=f"{d:%d}",
              company=clean(company), position=clean(position))
    try:
        name = S["folder_pattern"].format(**kw)
    except (KeyError, ValueError, IndexError):
        name = PATTERN_PRESETS[0].format(**kw)
    return clean(name, 120)


def read_rows():
    wb = load_workbook(find_excel())
    ws = sheet(wb)
    headers = [(c.column, str(c.value).strip()) for c in ws[1] if c.value is not None]
    rows = []
    for r in range(2, ws.max_row + 1):
        vals = {col: as_text(ws.cell(r, col).value) for col, _ in headers}
        if any(vals.values()):
            rows.append((r, vals))
    return headers, rows


def set_cell(cell, header, val):
    f = field_of(header)
    if f in DATE_FIELDS and val and parse_date(val):
        cell.value = parse_date(val)
        cell.number_format = "DD.MM.YYYY"
    elif f == "link":
        cell.value = val or None
        cell.hyperlink = val or None
        if val:
            cell.style = "Hyperlink"
    elif f == "folder":
        cell.value = val or None
        cell.hyperlink = os.path.join(job_dir(), val) if val else None
        if val:
            cell.style = "Hyperlink"
    else:
        cell.value = val if val != "" else None


def _save(wb, path):
    shutil.copy2(path, path + ".bak")
    wb.save(path)


def update_row(row, changes, headers):
    path = find_excel()
    ensure_closed(path)
    wb = load_workbook(path)
    ws = sheet(wb)
    names = dict(headers)
    for col, val in changes.items():
        set_cell(ws.cell(row, col), names[col], val)
    _save(wb, path)


def append_row(colvals, headers):
    path = find_excel()
    ensure_closed(path)
    wb = load_workbook(path)
    ws = sheet(wb)
    names = dict(headers)

    key_col = find_field(headers, "company") or headers[0][0]
    last = 1
    for r in range(ws.max_row, 1, -1):
        if ws.cell(r, key_col).value not in (None, ""):
            last = r
            break
    row = last + 1

    for col, h in headers:
        cell = ws.cell(row, col)
        if last > 1:
            above = ws.cell(last, col)
            if above.has_style:
                cell.font = copy(above.font)
                cell.border = copy(above.border)
                cell.fill = copy(above.fill)
                cell.alignment = copy(above.alignment)
                cell.number_format = above.number_format
        set_cell(cell, h, colvals.get(col, ""))

    for tbl in ws.tables.values():
        start, end = tbl.ref.split(":")
        end_col = re.match(r"[A-Z]+", end).group()
        end_row = int(re.search(r"\d+", end).group())
        if row > end_row:
            tbl.ref = f"{start}:{end_col}{row}"
            if tbl.autoFilter is not None:
                tbl.autoFilter.ref = tbl.ref
    _save(wb, path)


def create_application(colvals, headers):
    """Fill automatic fields, append the row, create folder and copy templates."""
    colvals = dict(colvals)

    def col(f):
        return find_field(headers, f)

    def get(f):
        c = col(f)
        return colvals.get(c, "") if c else ""

    folder = make_folder_name(get("date"), get("company"), get("position"))
    autos = {"id": folder, "folder": folder, "status": S["status_options"][0],
             "result": S["rav_options"][0], "last": get("date")}
    for f, v in autos.items():
        c = col(f)
        if c and not colvals.get(c):
            colvals[c] = v
    folder = get("folder") or folder

    append_row(colvals, headers)

    path = os.path.join(job_dir(), folder)
    os.makedirs(path, exist_ok=True)
    for t in S["templates"]:
        if os.path.isfile(t):
            dst = os.path.join(path, os.path.basename(t))
            if not os.path.exists(dst):
                shutil.copy2(t, dst)
    return path


def compute_stats(headers, rows):
    today = dt.date.today()
    dc, sc, lc = (find_field(headers, f) for f in ("date", "status", "last"))
    month = 0
    followups = set()
    for r, v in rows:
        d = parse_date(v.get(dc, "")) if dc else None
        if d and d.year == today.year and d.month == today.month:
            month += 1
        status = v.get(sc, "").strip().lower() if sc else ""
        if status in ("applied", ""):
            ref = parse_date(v.get(lc, "")) if lc else None
            ref = ref or d
            if ref and (today - ref.date()).days >= int(S["followup_days"]):
                followups.add(r)
    return month, followups


def open_path(path):
    if sys.platform.startswith("win"):
        os.startfile(path)
    elif sys.platform == "darwin":
        subprocess.Popen(["open", path])
    else:
        subprocess.Popen(["xdg-open", path])


def stop_sound():
    if sys.platform.startswith("win"):
        try:
            import winsound
            winsound.PlaySound(None, 0)
        except Exception:
            pass


def click_sound():
    if S["sound"] and S.get("click_sounds", True):
        play_sound(os.path.join(ASSETS, "click.wav"))


def play_sound(path):
    if not path or not os.path.isfile(path):
        return
    try:
        if sys.platform.startswith("win"):
            import winsound
            winsound.PlaySound(path, winsound.SND_FILENAME | winsound.SND_ASYNC)
        elif sys.platform == "darwin":
            subprocess.Popen(["afplay", path])
        else:
            subprocess.Popen(["aplay", "-q", path])
    except Exception:
        pass


# =====================================================================
# Look and feel
# =====================================================================
ACCENT = ("#6C5CE7", "#7F70F5")
ACCENT_HOVER = ("#5A4BD6", "#6C5CE7")
CARD = ("#FFFFFF", "#24242B")
BG = ("#F3F3F8", "#17171C")
MUTED = ("#6B6B7B", "#9A9AAE")
DANGER = ("#E5484D", "#F2555A")
SUBTLE_BTN = ("#E6E6EF", "#2F2F38")
SUBTLE_HOVER = ("#D8D8E4", "#3A3A45")
TEXT = ("#1A1A22", "#EDEDF3")
TRACK = ("#E3E3EC", "#33333D")
CONFETTI = ["#6C5CE7", "#00B894", "#FDCB6E", "#E17055", "#0984E3",
            "#E84393", "#55EFC4", "#FAB1A0", "#A29BFE", "#FFEAA7"]

QUESTIONS = [
    # field, label, question, required
    ("company",  "Company",            "Which company?",                 True),
    ("position", "Position",           "Which position?",                True),
    ("link",     "Job ad link",        "Link to the job ad?",            False),
    ("date",     "Date applied",       "Date applied (dd.mm.yyyy)?",     True),
    ("workload", "Workload",           "Workload?",                      False),
    ("method",   "Application method", "How did you apply?",             False),
    ("source",   "Source",             "Where did you find it?",         False),
    ("location", "Location",           "Location?",                      False),
    ("address",  "Address",            "Company address?",               False),
    ("contact",  "Contact person",     "Contact person?",                False),
    ("email",    "Contact email",      "Contact email or phone?",        False),
    ("notes",    "Notes",              "Anything to note?",              False),
]
RAV_FIELDS = ("address", "contact", "email")


def F(size=14, weight="normal"):
    return ctk.CTkFont(family="Segoe UI", size=size, weight=weight)


def option_like(prefix, options, default):
    for o in options:
        if o.lower().startswith(prefix):
            return o
    return default


def today_str():
    return dt.date.today().strftime(DATE_FMT)


# =====================================================================
# Celebrations
# =====================================================================
GOLD = ["#FFD700", "#FFC300", "#F5B700", "#FFE066", "#E6A800", "#FFF1A8", "#D4AF37"]

GOAL_MESSAGES = [
    "WOW, YOU'RE AMAZING!",
    "HOTSHOT!",
    "UNSTOPPABLE!",
    "ON FIRE!",
    "APPLICATION MACHINE!",
    "LEGENDARY!",
    "RECRUITERS FEAR YOU",
    "BEAST MODE",
    "GODLIKE!",
    "DOUBLE DIGITS, BABY!",
    "ABSOLUTE UNIT",
    "HR'S WORST NIGHTMARE",
    "THE CV WHISPERER",
    "MONSTER OUTPUT",
    "HALL OF FAME",
    "LINKEDIN LEGEND",
    "YOUR RAV ADVISOR IS CRYING",
    "IT'S OVER 9000!",
    "COSMIC LEVEL",
    "YOU BROKE THE SCALE",
]


def goal_preset(k):
    """k = applications over the goal (0 = goal just reached)."""
    if k <= 0:
        return dict(text="YOU ROCK!", sub="Monthly goal reached")
    lvl = min(k, 12)
    msg = GOAL_MESSAGES[k - 1] if k <= len(GOAL_MESSAGES) else f"STILL GOING! +{k}"
    return dict(text=msg, sub=f"+{k} over your goal",
                rain=170 + 25 * lvl, duration=3.2 + 0.3 * lvl,
                cannons=lvl // 2, fireworks=max(0, lvl - 3),
                sparkles=8 * lvl if lvl >= 3 else 0, rings=1 if lvl >= 7 else 0,
                rainbow=lvl >= 6, wobble=lvl >= 5, shake=10 if lvl >= 8 else 0,
                palette=CONFETTI + GOLD if lvl >= 10 else CONFETTI)


def interview_preset(company):
    return dict(text="INTERVIEW!", sub=f"{company} wants to meet you", palette=GOLD,
                duration=5.5, rain=280, cannons=3, sparkles=60, rings=2, wobble=True,
                text_color="#FFC300", sound="sound_interview")


def offer_preset(company):
    return dict(text="JOB OFFER!!!", sub=f"{company} wants YOU", palette=GOLD + CONFETTI,
                duration=9.5, rain=420, cannons=8, fireworks=16, sparkles=140, rings=5,
                rainbow=True, wobble=True, shake=45, sound="sound_offer")


class Celebration:
    KEY = "#010203"
    active = None

    @classmethod
    def stop_active(cls, _e=None):
        if cls.active is not None:
            cls.active.stop()

    def __init__(self, root, text="YOU ROCK!", sub="", palette=None, duration=3.2, rain=170,
                 cannons=0, fireworks=0, sparkles=0, rings=0, rainbow=False, wobble=False,
                 shake=0, text_color=None, sound="sound_file", sound_path=None, on_done=None):
        import tkinter.font as tkfont
        Celebration.stop_active()
        Celebration.active = self
        self.after_id = None
        self.stopped = False
        root.update_idletasks()
        self.root = root
        w, h = root.winfo_width(), root.winfo_height()
        self.cx0, self.cy0 = root.winfo_rootx(), root.winfo_rooty()
        self.rx0, self.ry0 = root.winfo_x(), root.winfo_y()
        self.top = tk.Toplevel(root)
        self.top.overrideredirect(True)
        self.top.attributes("-topmost", True)
        self.top.geometry(f"{w}x{h}+{self.cx0}+{self.cy0}")
        bg = self.KEY
        if sys.platform.startswith("win"):
            try:
                self.top.attributes("-transparentcolor", self.KEY)
            except tk.TclError:
                bg = "#17171C"
        else:
            bg = "#17171C" if ctk.get_appearance_mode() == "Dark" else "#FFFFFF"
        self.cv = tk.Canvas(self.top, width=w, height=h, bg=bg, highlightthickness=0)
        self.cv.pack(fill="both", expand=True)
        self.w, self.h = w, h
        self.palette = palette or CONFETTI
        self.total = int(duration * 1000 / 16)
        self.frame = 0
        self.rainbow, self.wobble, self.shake = rainbow, wobble, shake
        self.text_color = text_color
        self.on_done = on_done
        self.parts = []

        for _ in range(rain):
            self.add_rain(random.uniform(-h, -10))
        self.events = []
        for i in range(cannons):
            self.events.append((4 + i * int(self.total * 0.7 / max(1, cannons)), self.cannon_burst))
        for _ in range(fireworks):
            self.events.append((12 + random.randint(0, int(self.total * 0.75)), self.firework))
        for i in range(rings):
            self.events.append((i * 14, self.ring))
        for _ in range(sparkles):
            self.parts.append(dict(kind="spark", x=random.uniform(0, w), y=random.uniform(0, h),
                                   r=random.uniform(4, 11), ph=random.uniform(0, 6.28),
                                   sp=random.uniform(0.08, 0.2),
                                   id=self.cv.create_polygon(0, 0, 0, 0, 0, 0,
                                                             fill=random.choice(GOLD + ["#FFFFFF"]),
                                                             outline="")))

        # text sized to fit the window
        self.base = min(w / 7, 84)
        f = tkfont.Font(family="Segoe UI Black", size=int(self.base), weight="bold")
        tw = f.measure(text)
        if tw > 0.86 * w:
            self.base *= 0.86 * w / tw
        self.text = text
        self.sub = sub
        self.shadow_id = self.cv.create_text(w / 2 + 3, h / 2 + 4, text="", fill="#2D2D3A")
        self.text_id = self.cv.create_text(w / 2, h / 2, text="", fill="#6C5CE7")
        dark = ctk.get_appearance_mode() == "Dark"
        self.sub_shadow = self.cv.create_text(w / 2 + 2, h / 2 + 2, text="",
                                              fill="#000000" if dark else "#FFFFFF")
        self.sub_id = self.cv.create_text(w / 2, h / 2, text="",
                                          fill="#FFFFFF" if dark else "#1A1A22")

        if sound_path is None:
            sound_path = script_rel(S.get(sound, "")) if S["sound"] else ""
        play_sound(sound_path)
        self.tick()

    # ---------- particles ----------
    def poly(self):
        return self.cv.create_polygon(0, 0, 0, 0, 0, 0, 0, 0,
                                      fill=random.choice(self.palette), outline="")

    def add_rain(self, y):
        self.parts.append(dict(kind="rain", x=random.uniform(0, self.w), y=y,
                               vx=random.uniform(-1.2, 1.2), vy=random.uniform(2.5, 6),
                               a=random.uniform(0, 6.28), spin=random.uniform(-0.25, 0.25),
                               sw=random.uniform(5, 10), sh=random.uniform(9, 16),
                               sway=random.uniform(0, 6.28), id=self.poly()))

    def cannon_burst(self):
        sx, sy = self.w / 860, self.h / 680
        for side in (0, 1):
            for _ in range(45):
                self.parts.append(dict(kind="cannon", x=self.w * side, y=self.h,
                                       vx=random.uniform(4, 11) * sx * (1 if side == 0 else -1),
                                       vy=-random.uniform(11, 19) * sy,
                                       a=random.uniform(0, 6.28), spin=random.uniform(-0.3, 0.3),
                                       sw=random.uniform(5, 10), sh=random.uniform(9, 16),
                                       sway=random.uniform(0, 6.28), id=self.poly()))

    def firework(self):
        import math
        x = random.uniform(0.12, 0.88) * self.w
        y = random.uniform(0.1, 0.45) * self.h
        col = random.choice(self.palette)
        for _ in range(46):
            ang = random.uniform(0, 6.28)
            sp = random.uniform(2.5, 7)
            self.parts.append(dict(kind="spark_burst", x=x, y=y, vx=math.cos(ang) * sp,
                                   vy=math.sin(ang) * sp, life=55, max=55,
                                   id=self.cv.create_oval(0, 0, 0, 0, fill=col, outline="")))

    def ring(self):
        self.parts.append(dict(kind="ring", r=10, life=40,
                               id=self.cv.create_oval(0, 0, 0, 0, outline=random.choice(GOLD), width=10)))

    def draw_flake(self, p):
        import math
        ca, sa = math.cos(p["a"]), math.sin(p["a"])
        sw = p["sw"] * abs(math.cos(p["sway"] * 1.7)) + 1
        pts = []
        for dx, dy in ((-sw, -p["sh"]), (sw, -p["sh"]), (sw, p["sh"]), (-sw, p["sh"])):
            pts += [p["x"] + dx * ca / 2 - dy * sa / 2, p["y"] + dx * sa / 2 + dy * ca / 2]
        self.cv.coords(p["id"], *pts)

    # ---------- main loop ----------
    def tick(self):
        import math
        import colorsys
        f = self.frame
        for fr, fn in self.events:
            if fr == f:
                fn()
        alive = []
        for p in self.parts:
            k = p["kind"]
            if k == "rain":
                p["sway"] += 0.08
                p["x"] += p["vx"] + math.sin(p["sway"]) * 0.8
                p["y"] += p["vy"]
                p["a"] += p["spin"]
                if p["y"] > self.h + 20:
                    if f < self.total - 70:
                        p["y"], p["x"] = -20, random.uniform(0, self.w)
                    else:
                        self.cv.delete(p["id"])
                        continue
                self.draw_flake(p)
            elif k == "cannon":
                p["vy"] += 0.32
                p["vx"] *= 0.985
                p["sway"] += 0.1
                p["x"] += p["vx"]
                p["y"] += min(p["vy"], 4.5)
                p["a"] += p["spin"]
                if p["y"] > self.h + 30:
                    self.cv.delete(p["id"])
                    continue
                self.draw_flake(p)
            elif k == "spark_burst":
                p["life"] -= 1
                if p["life"] <= 0:
                    self.cv.delete(p["id"])
                    continue
                p["vx"] *= 0.96
                p["vy"] = p["vy"] * 0.96 + 0.07
                p["x"] += p["vx"]
                p["y"] += p["vy"]
                r = 3.5 * p["life"] / p["max"] + 0.5
                self.cv.coords(p["id"], p["x"] - r, p["y"] - r, p["x"] + r, p["y"] + r)
            elif k == "ring":
                p["life"] -= 1
                if p["life"] <= 0:
                    self.cv.delete(p["id"])
                    continue
                p["r"] += max(self.w, self.h) / 45
                r = p["r"]
                self.cv.coords(p["id"], self.w / 2 - r, self.h / 2 - r, self.w / 2 + r, self.h / 2 + r)
                self.cv.itemconfig(p["id"], width=max(1, int(10 * p["life"] / 40)))
            elif k == "spark":
                s = p["r"] * max(0.0, math.sin(p["ph"] + f * p["sp"]))
                if f > self.total - 20:
                    s *= (self.total - f) / 20
                x, y = p["x"], p["y"]
                q = s * 0.28
                self.cv.coords(p["id"], x, y - s, x + q, y - q, x + s, y, x + q, y + q,
                               x, y + s, x - q, y + q, x - s, y, x - q, y - q)
            alive.append(p)
        self.parts = alive

        # text: pop in with overshoot, pulse, shrink away
        if f < 22:
            scale = 1.25 * math.sin(f / 22 * math.pi / 2) ** 2
        elif f < 30:
            scale = 1.25 - (f - 22) / 8 * 0.25
        elif f < self.total - 25:
            scale = 1.0 + 0.04 * math.sin((f - 30) / 6)
        else:
            scale = max(0.01, (self.total - f) / 25)
        size = max(1, int(self.base * scale))
        angle = 5 * math.sin(f / 8) if self.wobble else 0
        jx = jy = 0
        if f < self.shake:
            jx, jy = random.randint(-9, 9), random.randint(-9, 9)
            self.root.geometry(f"+{self.rx0 + jx}+{self.ry0 + jy}")
            self.top.geometry(f"+{self.cx0 + jx}+{self.cy0 + jy}")
        elif f == self.shake and self.shake:
            self.root.geometry(f"+{self.rx0}+{self.ry0}")
            self.top.geometry(f"+{self.cx0}+{self.cy0}")
        cx, cy = self.w / 2, self.h / 2 - (self.base * 0.55 if self.sub else 0)
        font = ("Segoe UI Black", size, "bold")
        self.cv.itemconfig(self.shadow_id, text=self.text, font=font, angle=angle)
        self.cv.coords(self.shadow_id, cx + 3, cy + 4)
        self.cv.itemconfig(self.text_id, text=self.text, font=font, angle=angle)
        self.cv.coords(self.text_id, cx, cy)
        if self.rainbow:
            r, g, b = colorsys.hsv_to_rgb((f * 0.02) % 1, 0.75, 1)
            col = f"#{int(r * 255):02x}{int(g * 255):02x}{int(b * 255):02x}"
        elif self.text_color:
            col = self.text_color
        else:
            col = CONFETTI[(f // 6) % 5]
        self.cv.itemconfig(self.text_id, fill=col)
        if self.sub and f > 14:
            ss = max(1, int(max(16, self.base * 0.34) * min(1.0, scale)))
            sy = cy + self.base * 1.05 + ss * 0.9
            sfont = ("Segoe UI", ss, "bold")
            self.cv.itemconfig(self.sub_shadow, text=self.sub, font=sfont)
            self.cv.coords(self.sub_shadow, self.w / 2 + 2, sy + 2)
            self.cv.itemconfig(self.sub_id, text=self.sub, font=sfont)
            self.cv.coords(self.sub_id, self.w / 2, sy)
        self.cv.tag_raise(self.shadow_id)
        self.cv.tag_raise(self.text_id)
        self.cv.tag_raise(self.sub_shadow)
        self.cv.tag_raise(self.sub_id)

        self.frame += 1
        if self.frame < self.total:
            self.after_id = self.top.after(16, self.tick)
        else:
            self.finish()
            if self.on_done:
                self.root.after(250, self.on_done)

    def finish(self):
        if self.shake:
            try:
                self.root.geometry(f"+{self.rx0}+{self.ry0}")
            except tk.TclError:
                pass
        try:
            self.top.destroy()
        except tk.TclError:
            pass
        if Celebration.active is self:
            Celebration.active = None

    def stop(self):
        """Interrupted by a click or key press: end now, no follow ups."""
        if self.stopped:
            return
        self.stopped = True
        if self.after_id:
            try:
                self.top.after_cancel(self.after_id)
            except tk.TclError:
                pass
        stop_sound()
        self.finish()


class GoalBar(tk.Canvas):
    """Progress bar that keeps growing past the goal: green up to the goal,
    a golden gradient for every bonus application, and a goal marker."""

    def __init__(self, parent, n, goal):
        dark = ctk.get_appearance_mode() == "Dark"
        self.dark = dark
        super().__init__(parent, height=44, highlightthickness=0, bd=0,
                         bg=CARD[1] if dark else CARD[0])
        self.n, self.goal = n, goal
        self.bind("<Configure>", lambda e: self.draw())

    @staticmethod
    def mix(c1, c2, t):
        a = [int(c1[i:i + 2], 16) for i in (1, 3, 5)]
        b = [int(c2[i:i + 2], 16) for i in (1, 3, 5)]
        return "#" + "".join(f"{int(x + (y - x) * t):02x}" for x, y in zip(a, b))

    def draw(self):
        self.delete("all")
        d = self.dark
        w = self.winfo_width()
        th, y, pad = 14, 14, 8
        x0, x1 = pad, w - pad
        n, goal = self.n, self.goal
        scale = max(n, goal)
        xg = x0 + (x1 - x0) * goal / scale
        xn = x0 + (x1 - x0) * min(n, scale) / scale
        track = TRACK[1] if d else TRACK[0]
        self.create_line(x0, y, x1, y, width=th, fill=track, capstyle="round")
        if n <= 0:
            pass
        elif n < goal:
            self.create_line(x0, y, xn, y, width=th, fill=ACCENT[1] if d else ACCENT[0], capstyle="round")
        else:
            self.create_line(x0, y, xg, y, width=th, fill="#00B894", capstyle="round")
        if n > goal:
            steps = 40
            for i in range(steps):
                a = xg + (xn - xg) * i / steps
                b = xg + (xn - xg) * (i + 1) / steps
                col = self.mix("#FF8C42", "#FFD700", i / (steps - 1))
                self.create_line(a, y, b + 1, y, width=th, fill=col,
                                 capstyle="round" if i == steps - 1 else "butt")
            # a thin divider for each bonus application
            card = CARD[1] if d else CARD[0]
            if n - goal <= 40:
                for k in range(1, n - goal):
                    xk = x0 + (x1 - x0) * (goal + k) / scale
                    self.create_line(xk, y - th / 2, xk, y + th / 2, fill=card, width=2)
            self.create_text(xn - 2, y - 1, text="★", fill="#FFFFFF", font=("Segoe UI", 9, "bold"),
                             anchor="e")
        # goal marker
        if 0 < goal and n >= goal and n > 0 and goal < scale:
            self.create_line(xg, y - th / 2 - 4, xg, y + th / 2 + 4, fill=TEXT[1] if d else TEXT[0], width=2)
        self.create_text(xg, y + th / 2 + 10, text=f"GOAL {goal}",
                         fill=MUTED[1] if d else MUTED[0], font=("Segoe UI", 8, "bold"),
                         anchor="e" if xg > x1 - 30 else "center")


# =====================================================================
# App
# =====================================================================
class App:
    def __init__(self, root):
        self.root = root
        root.title("Job applications")
        root.geometry("860x680")
        root.minsize(720, 560)
        root.configure(fg_color=BG)
        self.mode = None
        self.body = ctk.CTkFrame(root, fg_color="transparent")
        self.body.pack(fill="both", expand=True, padx=28, pady=24)
        root.bind("<Return>", self.on_enter)
        root.bind("<Escape>", self.on_escape)
        root.bind_all("<Button-1>", Celebration.stop_active, add="+")
        root.bind_all("<Key>", Celebration.stop_active, add="+")
        self.gear_img = None
        p = os.path.join(ASSETS, "icon_setting.png")
        if os.path.isfile(p):
            try:
                from PIL import Image
                img = Image.open(p)
                self.gear_img = ctk.CTkImage(light_image=img, dark_image=img, size=(24, 24))
            except Exception:
                self.gear_img = None
        self.home()

    # ---------- helpers ----------
    def clear(self, mode):
        self.mode = mode
        for w in self.body.winfo_children():
            w.destroy()

    def card(self, parent=None, **kw):
        return ctk.CTkFrame(parent or self.body, fg_color=CARD, corner_radius=16, **kw)

    def button(self, parent, text, cmd, kind="primary", **kw):
        styles = {
            "primary": dict(fg_color=ACCENT, hover_color=ACCENT_HOVER, text_color="white"),
            "subtle": dict(fg_color=SUBTLE_BTN, hover_color=SUBTLE_HOVER, text_color=TEXT),
            "danger": dict(fg_color="transparent", hover_color=SUBTLE_BTN, text_color=DANGER),
            "interview": dict(fg_color="#00B894", hover_color="#00A383", text_color="white"),
            "gold": dict(fg_color="#F5B700", hover_color="#E0A600", text_color="#3A2A00"),
        }
        kw.setdefault("height", 40)
        kw.setdefault("corner_radius", 12)
        kw.setdefault("font", F(14, "bold" if kind in ("primary", "interview", "gold") else "normal"))
        return ctk.CTkButton(parent, text=text, command=self.clicky(cmd), **styles[kind], **kw)

    def clicky(self, cmd):
        def run(*a):
            click_sound()
            return cmd(*a)
        return run

    def header(self, title, back=None, gear=False):
        bar = ctk.CTkFrame(self.body, fg_color="transparent")
        bar.pack(fill="x", pady=(0, 16))
        if back:
            self.button(bar, "‹  Back", back, "subtle", width=90, height=34).pack(side="left", padx=(0, 14))
        ctk.CTkLabel(bar, text=title, font=F(26, "bold"), text_color=TEXT).pack(side="left")
        if gear:
            ctk.CTkButton(bar, text="" if self.gear_img else "⚙", image=self.gear_img,
                          width=40, height=40, corner_radius=12, font=F(20),
                          fg_color="transparent", hover_color=SUBTLE_BTN, text_color=TEXT,
                          command=self.clicky(self.settings_view)).pack(side="right")
            on = S["sound"]
            ctk.CTkButton(bar, text="♪  Sound on" if on else "♪  Muted", width=110, height=34,
                          corner_radius=17, font=F(13, "bold" if on else "normal"),
                          fg_color=SUBTLE_BTN if on else "transparent", hover_color=SUBTLE_HOVER,
                          border_width=0 if on else 1, border_color=MUTED,
                          text_color=TEXT if on else MUTED,
                          command=self.toggle_mute).pack(side="right", padx=(0, 8))
        return bar

    def toggle_mute(self):
        S["sound"] = not S["sound"]
        if not S["sound"]:
            stop_sound()
        try:
            save_settings()
        except Exception:
            pass
        click_sound()
        self.home()

    def toast(self, text):
        lbl = ctk.CTkLabel(self.root, text=f"  {text}  ", font=F(13, "bold"),
                           fg_color=("#1A1A22", "#EDEDF3"), text_color=("#FFFFFF", "#1A1A22"),
                           corner_radius=12, height=36)
        lbl.place(relx=0.5, rely=0.94, anchor="center")
        self.root.after(2400, lbl.destroy)

    def error(self, e):
        messagebox.showerror("Something went wrong", str(e), parent=self.root)

    def on_enter(self, _e=None):
        actions = {"q": self.q_next, "summary": self.save_flow, "list": self.open_selected,
                   "form": self.save_form, "export": self.do_export,
                   "pick": self.confirm_pick}
        if self.mode in actions:
            actions[self.mode]()

    def on_escape(self, _e=None):
        if self.mode == "home":
            self.root.destroy()
        elif self.mode in ("q", "summary"):
            self.discard_flow()
        elif self.mode == "form":
            self.exit_form()
        else:
            self.home()

    def load(self):
        try:
            self.headers, self.rows = read_rows()
            return True
        except Exception as e:
            self.error(e)
            return False

    def month_count(self):
        try:
            return compute_stats(self.headers, self.rows)[0]
        except Exception:
            return 0

    def celebrate(self, preset):
        self.root.after(300, lambda: Celebration(self.root, **preset))

    def maybe_celebrate(self, before):
        """Celebrate when a save reached or went beyond the monthly goal."""
        try:
            headers, rows = read_rows()
        except Exception:
            return
        n, _ = compute_stats(headers, rows)
        goal = int(S["monthly_goal"] or 0)
        if goal > 0 and n > before and n >= goal:
            self.celebrate(goal_preset(n - goal))

    # =================================================================
    # Home
    # =================================================================
    def home(self):
        self.clear("home")
        self.header("Job applications", gear=True)
        try:
            headers, rows = read_rows()
            n, follow = compute_stats(headers, rows)
            ok = True
        except Exception as e:
            ok, n, follow, err = False, 0, set(), e

        goal = int(S["monthly_goal"] or 0)
        c = self.card()
        c.pack(fill="x", pady=(0, 14))
        inner = ctk.CTkFrame(c, fg_color="transparent")
        inner.pack(fill="x", padx=24, pady=20)
        month_name = dt.date.today().strftime("%B %Y")
        ctk.CTkLabel(inner, text=month_name.upper(), font=F(12, "bold"),
                     text_color=MUTED).pack(anchor="w")
        if ok:
            over = n - goal if goal else 0
            row = ctk.CTkFrame(inner, fg_color="transparent")
            row.pack(fill="x", pady=(2, 8))
            ctk.CTkLabel(row, text=str(n), font=F(48, "bold"),
                         text_color=("#D99A00", "#FFC300") if over > 0 else TEXT).pack(side="left")
            ctk.CTkLabel(row, text=f"  /  {goal} applications" if goal else "  applications",
                         font=F(18), text_color=MUTED).pack(side="left", pady=(18, 0))
            if over > 0:
                tier = GOAL_MESSAGES[over - 1] if over <= len(GOAL_MESSAGES) else f"STILL GOING! +{over}"
                ctk.CTkLabel(row, text=f"  +{over}   {tier}  ", font=F(13, "bold"),
                             fg_color="#F5B700", text_color="#3A2A00", corner_radius=14,
                             height=30).pack(side="right", pady=(14, 0))
            elif goal and n == goal:
                ctk.CTkLabel(row, text="  ✓  GOAL REACHED  ", font=F(13, "bold"),
                             fg_color="#00B894", text_color="white", corner_radius=14,
                             height=30).pack(side="right", pady=(14, 0))
            if goal:
                GoalBar(inner, n, goal).pack(fill="x")
                if n < goal:
                    msg = f"{goal - n} to go this month"
                elif over == 0:
                    msg = f"Everything from here is bonus. Next unlock: {GOAL_MESSAGES[0]}"
                else:
                    nxt = GOAL_MESSAGES[over] if over < len(GOAL_MESSAGES) else f"STILL GOING! +{over + 1}"
                    msg = (f"Goal smashed with {over} bonus application{'s' if over > 1 else ''}. "
                           f"Next unlock: {nxt}")
                ctk.CTkLabel(inner, text=msg, font=F(13), text_color=MUTED).pack(anchor="w", pady=(6, 0))
        else:
            ctk.CTkLabel(inner, text=f"Could not read the tracker:\n{err}", font=F(13),
                         text_color=DANGER, justify="left").pack(anchor="w", pady=8)

        if follow:
            fc = self.card()
            fc.pack(fill="x", pady=(0, 14))
            fi = ctk.CTkFrame(fc, fg_color="transparent")
            fi.pack(fill="x", padx=24, pady=14)
            k = len(follow)
            ctk.CTkLabel(fi, text=f"⏰  {k} application{'s' if k > 1 else ''} without news "
                                  f"for {S['followup_days']}+ days",
                         font=F(14), text_color=TEXT).pack(side="left")
            self.button(fi, "Review", lambda: self.list_view(only_follow=True), "subtle",
                        width=100, height=34).pack(side="right")

        grid = ctk.CTkFrame(self.body, fg_color="transparent")
        grid.pack(fill="x", pady=(6, 0))
        grid.columnconfigure((0, 1), weight=1, uniform="a")
        self.button(grid, "＋  New entry", self.new_flow, height=56,
                    font=F(16, "bold")).grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 10))
        self.button(grid, "★  Mark interview received", lambda: self.list_view(pick="interview"),
                    "interview", height=48).grid(row=1, column=0, sticky="ew", padx=(0, 5), pady=(0, 10))
        self.button(grid, "♛  Mark job offer", lambda: self.list_view(pick="offer"),
                    "gold", height=48).grid(row=1, column=1, sticky="ew", padx=(5, 0), pady=(0, 10))
        self.button(grid, "Edit entries", self.list_view, "subtle", height=48).grid(
            row=2, column=0, sticky="ew", padx=(0, 5))
        self.button(grid, "RAV export", self.export_view, "subtle", height=48).grid(
            row=2, column=1, sticky="ew", padx=(5, 0))
        self.button(grid, "Exit", self.root.destroy, "danger", height=40).grid(
            row=3, column=0, columnspan=2, sticky="ew", pady=(10, 0))

    # =================================================================
    # New entry: question flow
    # =================================================================
    def new_flow(self):
        if not self.load():
            return
        self.qi = 0
        self.reviewing = False
        d = S["defaults"]
        self.answers = {"date": today_str(), "workload": d.get("workload", ""),
                        "method": d.get("method", ""), "source": d.get("source", ""),
                        "location": d.get("location", "")}
        self.show_q()

    def show_q(self):
        self.clear("q")
        field, label, q, req = QUESTIONS[self.qi]
        self.header("New entry")
        c = self.card()
        c.pack(fill="x")
        inner = ctk.CTkFrame(c, fg_color="transparent")
        inner.pack(fill="x", padx=28, pady=26)
        ctk.CTkLabel(inner, text=f"QUESTION {self.qi + 1} OF {len(QUESTIONS)}",
                     font=F(12, "bold"), text_color=MUTED).pack(anchor="w")
        pb = ctk.CTkProgressBar(inner, height=6, progress_color=ACCENT, fg_color=TRACK)
        pb.pack(fill="x", pady=(6, 18))
        pb.set((self.qi + 1) / len(QUESTIONS))
        ctk.CTkLabel(inner, text=q + ("  *" if req else ""), font=F(22, "bold"),
                     text_color=TEXT).pack(anchor="w", pady=(0, 10))
        self.entry = ctk.CTkEntry(inner, height=46, corner_radius=12, font=F(16),
                                  border_color=ACCENT)
        self.entry.pack(fill="x")
        self.entry.insert(0, self.answers.get(field, ""))
        self.root.after(60, self.focus_entry)
        ctk.CTkLabel(inner, text="Enter  next     Esc  cancel", font=F(12),
                     text_color=MUTED).pack(anchor="w", pady=(8, 0))

        btns = ctk.CTkFrame(self.body, fg_color="transparent")
        btns.pack(fill="x", pady=16)
        self.button(btns, "‹  Back", self.q_back, "subtle", width=120).pack(side="left")
        last = self.reviewing or self.qi == len(QUESTIONS) - 1
        self.button(btns, "Review  ›" if last else "Next  ›", self.q_next, width=140).pack(side="right")

    def focus_entry(self):
        if self.mode == "q" and self.entry.winfo_exists():
            self.entry.focus_set()
            self.entry.select_range(0, "end")

    def q_next(self):
        field, label, _, req = QUESTIONS[self.qi]
        val = self.entry.get().strip()
        if req and not val:
            self.toast(f"{label} is required")
            return
        if field == "date" and not parse_date(val):
            self.toast("Use the format dd.mm.yyyy")
            return
        self.answers[field] = val
        if self.reviewing or self.qi == len(QUESTIONS) - 1:
            self.summary()
        else:
            self.qi += 1
            self.show_q()

    def q_back(self):
        if self.qi == 0:
            self.discard_flow()
            return
        self.answers[QUESTIONS[self.qi][0]] = self.entry.get().strip()
        self.qi -= 1
        self.show_q()

    def discard_flow(self):
        if messagebox.askyesno("Discard", "Discard this application?", parent=self.root):
            self.home()

    def summary(self):
        self.clear("summary")
        self.reviewing = True
        self.header("Check your entries")
        sc = ctk.CTkScrollableFrame(self.body, fg_color=CARD, corner_radius=16)
        sc.pack(fill="both", expand=True)
        sc.columnconfigure(1, weight=1)
        r = 0
        for idx, (field, label, _, _) in enumerate(QUESTIONS):
            val = self.answers.get(field, "")
            missing = field in RAV_FIELDS and not val
            ctk.CTkLabel(sc, text=label, font=F(13, "bold"), text_color=MUTED).grid(
                row=r, column=0, padx=(18, 14), pady=5, sticky="nw")
            ctk.CTkLabel(sc, text=val or ("missing for RAV" if missing else "·"),
                         font=F(14), wraplength=430, justify="left",
                         text_color=DANGER if missing else TEXT).grid(row=r, column=1, pady=5, sticky="w")
            ctk.CTkButton(sc, text="Edit", width=56, height=28, corner_radius=8, font=F(12),
                          fg_color=SUBTLE_BTN, hover_color=SUBTLE_HOVER, text_color=TEXT,
                          command=self.clicky(lambda i=idx: self.jump(i))).grid(row=r, column=2, padx=14, pady=5)
            r += 1
        folder = make_folder_name(self.answers["date"], self.answers["company"], self.answers["position"])
        ctk.CTkLabel(sc, text="AUTOMATIC", font=F(11, "bold"), text_color=MUTED).grid(
            row=r, column=0, padx=18, pady=(14, 2), sticky="w")
        r += 1
        for label, val in (("ID and folder", folder), ("Status", S["status_options"][0]),
                           ("Last update", self.answers["date"]), ("Result for RAV", S["rav_options"][0])):
            ctk.CTkLabel(sc, text=label, font=F(13, "bold"), text_color=MUTED).grid(
                row=r, column=0, padx=(18, 14), pady=3, sticky="w")
            ctk.CTkLabel(sc, text=val, font=F(14), text_color=MUTED).grid(row=r, column=1, pady=3, sticky="w")
            r += 1
        btns = ctk.CTkFrame(self.body, fg_color="transparent")
        btns.pack(fill="x", pady=(16, 0))
        self.button(btns, "Discard", self.discard_flow, "danger", width=120).pack(side="left")
        self.button(btns, "Save  ✓", self.save_flow, width=160).pack(side="right")
        self.root.focus_set()

    def jump(self, i):
        self.qi = i
        self.show_q()

    def save_flow(self):
        before = self.month_count()
        colvals = {}
        for field, val in self.answers.items():
            c = find_field(self.headers, field)
            if c:
                colvals[c] = val
        try:
            path = create_application(colvals, self.headers)
        except Exception as e:
            self.error(e)
            return
        open_path(path)
        self.home()
        self.toast(f"{self.answers['company']} saved")
        self.maybe_celebrate(before)

    # =================================================================
    # Edit: list
    # =================================================================
    def style_tree(self):
        dark = ctk.get_appearance_mode() == "Dark"
        bg = CARD[1] if dark else CARD[0]
        fg = TEXT[1] if dark else TEXT[0]
        head = "#2F2F38" if dark else "#ECECF4"
        st = ttk.Style()
        st.theme_use("clam")
        st.configure("App.Treeview", background=bg, fieldbackground=bg, foreground=fg,
                     rowheight=34, borderwidth=0, font=("Segoe UI", 11))
        st.configure("App.Treeview.Heading", background=head, foreground=fg, relief="flat",
                     font=("Segoe UI", 11, "bold"), padding=8)
        st.map("App.Treeview", background=[("selected", ACCENT[1] if dark else ACCENT[0])],
               foreground=[("selected", "white")])
        st.map("App.Treeview.Heading", background=[("active", head)])
        st.layout("App.Treeview", [("App.Treeview.treearea", {"sticky": "nswe"})])
        return dark

    def list_view(self, only_follow=False, pick=None):
        if not self.load():
            return
        self.pick = pick
        self.clear("pick" if pick else "list")
        if pick:
            title = "Interview received" if pick == "interview" else "Job offer received"
            self.header(title, back=self.home)
            ctk.CTkLabel(self.body, text="Which application? Pick it and confirm.", font=F(14),
                         text_color=MUTED).pack(anchor="w", pady=(0, 10))
        else:
            bar = self.header("Applications", back=self.home)
            self.button(bar, "＋  New entry", self.new_form, width=140, height=36).pack(side="right")

        tools = ctk.CTkFrame(self.body, fg_color="transparent")
        tools.pack(fill="x", pady=(0, 10))
        self.search = ctk.CTkEntry(tools, placeholder_text="Search company, position, status...",
                                   height=38, corner_radius=12, font=F(14))
        self.search.pack(side="left", fill="x", expand=True)
        self.search.bind("<KeyRelease>", lambda e: self.fill_tree())
        self.search.bind("<Down>", lambda e: self.tree.focus_set())
        self.filter = ctk.CTkSegmentedButton(tools, values=["Open", "All"] if pick else ["All", "Follow up"],
                                             height=38,
                                             font=F(13), selected_color=ACCENT,
                                             selected_hover_color=ACCENT_HOVER,
                                             fg_color=TRACK, unselected_color=TRACK,
                                             unselected_hover_color=SUBTLE_HOVER, text_color=TEXT,
                                             command=lambda v: self.fill_tree())
        self.filter.pack(side="right", padx=(10, 0))
        self.filter.set("Open" if pick else ("Follow up" if only_follow else "All"))

        dark = self.style_tree()
        c = self.card()
        c.pack(fill="both", expand=True)
        wanted = ("date", "company", "position", "status", "result")
        self.tree_cols = [(find_field(self.headers, f), f) for f in wanted if find_field(self.headers, f)]
        names = dict(self.headers)
        self.tree = ttk.Treeview(c, columns=[str(cl) for cl, _ in self.tree_cols],
                                 show="headings", style="App.Treeview", selectmode="browse")
        short = {"date": "Date", "company": "Company", "position": "Position",
                 "status": "Status", "result": "RAV"}
        for cl, f in self.tree_cols:
            wide = f in ("company", "position")
            self.tree.heading(str(cl), text=short[f], anchor="w")
            self.tree.column(str(cl), width=220 if wide else 100, minwidth=80,
                             stretch=wide, anchor="w")
        sb = ctk.CTkScrollbar(c, command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y", padx=(0, 6), pady=10)
        self.tree.pack(fill="both", expand=True, padx=(12, 0), pady=12)
        self.tree.bind("<Double-1>", lambda e: self.confirm_pick() if self.pick else self.open_selected())
        alt = "#2A2A32" if dark else "#F7F7FB"
        self.tree.tag_configure("alt", background=alt)
        self.tree.tag_configure("interview", foreground="#00B894")
        self.tree.tag_configure("offer", foreground="#E1A800")
        self.tree.tag_configure("rejected", foreground="#8A8A99")

        _, self.follow = compute_stats(self.headers, self.rows)
        self.fill_tree()

        btns = ctk.CTkFrame(self.body, fg_color="transparent")
        btns.pack(fill="x", pady=(14, 0))
        if pick:
            label = "★  Mark as interview" if pick == "interview" else "♛  Mark as job offer"
            kind = "interview" if pick == "interview" else "gold"
            ctk.CTkLabel(btns, text="Double click or Enter to confirm", font=F(12),
                         text_color=MUTED).pack(side="left")
            self.button(btns, label, self.confirm_pick, kind, width=200).pack(side="right")
        else:
            ctk.CTkLabel(btns, text="Double click or Enter to edit", font=F(12),
                         text_color=MUTED).pack(side="left")
            self.button(btns, "Edit selected", self.open_selected, width=150).pack(side="right")
        self.root.after(60, lambda: self.mode in ("list", "pick") and self.search.winfo_exists()
                        and self.search.focus_set())

    def fill_tree(self):
        q = self.search.get().lower().strip()
        fv = self.filter.get()
        only = fv == "Follow up"
        sc = find_field(self.headers, "status")
        self.tree.delete(*self.tree.get_children())
        dc = find_field(self.headers, "date")
        ordered = sorted(self.rows, reverse=True,
                         key=lambda t: (parse_date(t[1].get(dc, "")) or dt.datetime.min, t[0]))
        i = 0
        for r, vals in ordered:
            if only and r not in self.follow:
                continue
            if fv == "Open" and not self.is_open(vals):
                continue
            if q and not any(q in v.lower() for v in vals.values()):
                continue
            tags = ["alt"] if i % 2 else []
            st = vals.get(sc, "").lower() if sc else ""
            if st.startswith("interview"):
                tags.append("interview")
            elif st.startswith("offer"):
                tags.append("offer")
            elif st.startswith("rejected"):
                tags.append("rejected")
            self.tree.insert("", "end", iid=str(r), tags=tags,
                             values=[vals[cl] for cl, _ in self.tree_cols])
            i += 1
        kids = self.tree.get_children()
        if kids:
            self.tree.selection_set(kids[0])
            self.tree.focus(kids[0])

    def is_open(self, vals):
        sc = find_field(self.headers, "status")
        st = vals.get(sc, "").lower() if sc else ""
        closed = ("interview", "offer", "rejected", "withdrawn") if self.pick == "interview" \
            else ("offer", "rejected", "withdrawn")
        return not st.startswith(closed)

    def confirm_pick(self):
        sel = self.tree.selection()
        if not sel:
            return
        row = int(sel[0])
        vals = dict(self.rows)[row]
        kind = self.pick
        sc, lc, rc, cc = (find_field(self.headers, f) for f in ("status", "last", "result", "company"))
        ch = {}
        if sc:
            ch[sc] = option_like(kind, S["status_options"], "Interview" if kind == "interview" else "Offer")
        if lc:
            ch[lc] = today_str()
        if rc and vals.get(rc, "").strip().lower() in ("", S["rav_options"][0].lower()):
            ch[rc] = option_like("interview", S["rav_options"], "Interview")
        try:
            update_row(row, ch, self.headers)
        except Exception as e:
            self.error(e)
            return
        company = vals.get(cc, "") if cc else ""
        self.home()
        self.celebrate(interview_preset(company) if kind == "interview" else offer_preset(company))

    def open_selected(self):
        sel = self.tree.selection()
        if sel:
            self.edit_form(int(sel[0]))

    # =================================================================
    # Edit / new: form with every column
    # =================================================================
    def build_form(self, title, values):
        self.clear("form")
        self.header(title, back=self.exit_form)
        sc = ctk.CTkScrollableFrame(self.body, fg_color=CARD, corner_radius=16)
        sc.pack(fill="both", expand=True)
        sc.columnconfigure(1, weight=1)
        self.fields = {}
        for i, (col, h) in enumerate(self.headers):
            f = field_of(h)
            ctk.CTkLabel(sc, text=h, font=F(13, "bold"), text_color=MUTED).grid(
                row=i, column=0, padx=(18, 14), pady=6, sticky="w")
            if f in ("status", "result"):
                opts = S["status_options"] if f == "status" else S["rav_options"]
                w = ctk.CTkComboBox(sc, values=opts, height=36, corner_radius=10, font=F(14),
                                    button_color=ACCENT, button_hover_color=ACCENT_HOVER)
                w.set(values.get(col, ""))
            else:
                ph = "automatic" if f in ("id", "folder") and self.form_new else ""
                w = ctk.CTkEntry(sc, height=36, corner_radius=10, font=F(14), placeholder_text=ph)
                if values.get(col):
                    w.insert(0, values[col])
            w.grid(row=i, column=1, padx=(0, 18), pady=6, sticky="ew")
            self.fields[col] = (w, values.get(col, ""), h)
        note = ("Empty automatic fields are filled in when you save." if self.form_new else
                "Last update is set to today when you change something.")
        btns = ctk.CTkFrame(self.body, fg_color="transparent")
        btns.pack(fill="x", pady=(14, 0))
        self.button(btns, "Exit", self.exit_form, "subtle", width=120).pack(side="left")
        ctk.CTkLabel(btns, text=note, font=F(12), text_color=MUTED).pack(side="left", padx=14)
        self.button(btns, "Save  ✓", self.save_form, width=150).pack(side="right")

    def edit_form(self, row):
        self.form_new = False
        self.form_row = row
        self.build_form("Edit entry", dict(self.rows)[row])

    def new_form(self):
        self.form_new = True
        vals = {}
        d = S["defaults"]
        prefill = {"date": today_str(), "last": today_str(),
                   "status": S["status_options"][0], "result": S["rav_options"][0],
                   "workload": d.get("workload", ""), "method": d.get("method", ""),
                   "source": d.get("source", ""), "location": d.get("location", "")}
        for f, v in prefill.items():
            c = find_field(self.headers, f)
            if c:
                vals[c] = v
        self.build_form("New entry", vals)

    def form_values(self):
        return {col: w.get().strip() for col, (w, _, _) in self.fields.items()}

    def form_changes(self):
        return {col: w.get().strip() for col, (w, orig, _) in self.fields.items()
                if w.get().strip() != orig}

    def exit_form(self):
        dirty = self.form_changes()
        if not dirty or messagebox.askyesno("Exit", "Exit without saving?", parent=self.root):
            self.list_view()

    def save_form(self):
        vals = self.form_values()
        for col, (_, _, h) in self.fields.items():
            if field_of(h) in DATE_FIELDS and vals[col] and not parse_date(vals[col]):
                self.toast(f"{h}: use dd.mm.yyyy")
                return
        if self.form_new:
            for f, label in (("company", "Company"), ("position", "Position"), ("date", "Date applied")):
                c = find_field(self.headers, f)
                if c and not vals.get(c):
                    self.toast(f"{label} is required")
                    return
            before = self.month_count()
            try:
                path = create_application(vals, self.headers)
            except Exception as e:
                self.error(e)
                return
            open_path(path)
            self.list_view()
            self.toast("Entry added")
            self.maybe_celebrate(before)
            return

        ch = self.form_changes()
        if not ch:
            self.list_view()
            return
        lc = find_field(self.headers, "last")
        if lc and lc not in ch:
            ch[lc] = today_str()
        try:
            update_row(self.form_row, ch, self.headers)
        except Exception as e:
            self.error(e)
            return
        self.list_view()
        self.toast("Changes saved")
        sc, cc = find_field(self.headers, "status"), find_field(self.headers, "company")
        if sc in ch:
            new, old = ch[sc].lower(), self.fields[sc][1].lower()
            company = vals.get(cc, "") if cc else ""
            if new.startswith("offer") and not old.startswith("offer"):
                self.celebrate(offer_preset(company))
            elif new.startswith("interview") and not old.startswith(("interview", "offer")):
                self.celebrate(interview_preset(company))

    # =================================================================
    # RAV export
    # =================================================================
    def export_view(self):
        if not self.load():
            return
        self.clear("export")
        self.header("RAV export", back=self.home)
        dc = find_field(self.headers, "date")
        months = sorted({(d.year, d.month) for _, v in self.rows
                         for d in [parse_date(v.get(dc, ""))] if d}, reverse=True)
        today = dt.date.today()
        if not months:
            months = [(today.year, today.month)]
        self.month_map = {f"{calendar.month_name[m]} {y}": (y, m) for y, m in months}
        labels = list(self.month_map)
        prev = (today.year - (today.month == 1), (today.month - 2) % 12 + 1)
        default = labels[0]
        if today.day <= 5 and prev in months:      # form is due on the 5th
            default = f"{calendar.month_name[prev[1]]} {prev[0]}"

        c = self.card()
        c.pack(fill="x")
        inner = ctk.CTkFrame(c, fg_color="transparent")
        inner.pack(fill="x", padx=28, pady=24)
        ctk.CTkLabel(inner, text="Month", font=F(13, "bold"), text_color=MUTED).pack(anchor="w")
        self.month_box = ctk.CTkComboBox(inner, values=labels, height=40, corner_radius=10,
                                         font=F(15), width=260, button_color=ACCENT,
                                         command=lambda v: self.export_info())
        self.month_box.pack(anchor="w", pady=(4, 12))
        self.month_box.set(default)
        self.info = ctk.CTkLabel(inner, text="", font=F(14), text_color=TEXT, justify="left")
        self.info.pack(anchor="w")
        ctk.CTkLabel(inner, text="Columns follow the RAV form: date, company, address, contact person, "
                                 "email / phone, position, workload, method, result.",
                     font=F(12), text_color=MUTED, wraplength=640, justify="left").pack(anchor="w", pady=(12, 0))
        btns = ctk.CTkFrame(self.body, fg_color="transparent")
        btns.pack(fill="x", pady=16)
        self.button(btns, "Export to Excel  ›", self.do_export, width=200).pack(side="right")
        self.export_info()

    def export_rows(self):
        y, m = self.month_map.get(self.month_box.get(), (0, 0))
        dc = find_field(self.headers, "date")
        out = []
        for _, v in self.rows:
            d = parse_date(v.get(dc, ""))
            if d and d.year == y and d.month == m:
                out.append((d, v))
        return sorted(out, key=lambda t: t[0])

    def export_info(self):
        rows = self.export_rows()
        missing = 0
        for _, v in rows:
            if any(not v.get(find_field(self.headers, f), "") for f in RAV_FIELDS
                   if find_field(self.headers, f)):
                missing += 1
        txt = f"{len(rows)} applications"
        if missing:
            txt += f"\n{missing} with missing address or contact details"
        self.info.configure(text=txt, text_color=DANGER if missing else TEXT)

    def do_export(self):
        rows = self.export_rows()
        if not rows:
            self.toast("No applications in that month")
            return
        y, m = self.month_map[self.month_box.get()]
        cols = [("Date", "date"), ("Company", "company"), ("Address", "address"),
                ("Contact person", "contact"), ("Email / phone", "email"),
                ("Position", "position"), ("Workload", "workload"),
                ("Application method", "method"), ("Result", "result")]
        wb = Workbook()
        ws = wb.active
        ws.title = f"{y}_{m:02d}"
        ws.append([c for c, _ in cols])
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for d, v in rows:
            line = []
            for _, f in cols:
                c = find_field(self.headers, f)
                line.append(d if f == "date" else (v.get(c, "") if c else ""))
            ws.append(line)
        for r in range(2, ws.max_row + 1):
            ws.cell(r, 1).number_format = "DD.MM.YYYY"
        for i, (name, _) in enumerate(cols, start=1):
            width = max([len(name)] + [len(str(ws.cell(r, i).value or "")) for r in range(2, ws.max_row + 1)])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(50, width + 3)
        out_dir = os.path.join(job_dir(), "RAV_exports")
        os.makedirs(out_dir, exist_ok=True)
        path = os.path.join(out_dir, f"RAV_{y}_{m:02d}.xlsx")
        try:
            ensure_closed(path)
            wb.save(path)
        except Exception as e:
            self.error(e)
            return
        open_path(path)
        self.toast("Export saved")

    # =================================================================
    # Settings
    # =================================================================
    def settings_view(self):
        self.clear("settings")
        self.header("Settings", back=self.home)
        sc = ctk.CTkScrollableFrame(self.body, fg_color="transparent")
        sc.pack(fill="both", expand=True)
        self.set_w = {}

        def section(title):
            c = self.card(sc)
            c.pack(fill="x", pady=(0, 12))
            ctk.CTkLabel(c, text=title.upper(), font=F(12, "bold"), text_color=MUTED).pack(
                anchor="w", padx=22, pady=(16, 6))
            body = ctk.CTkFrame(c, fg_color="transparent")
            body.pack(fill="x", padx=22, pady=(0, 16))
            body.columnconfigure(1, weight=1)
            return body

        def row(parent, r, label, widget, extra=None):
            ctk.CTkLabel(parent, text=label, font=F(14), text_color=TEXT).grid(
                row=r, column=0, sticky="w", pady=6, padx=(0, 16))
            widget.grid(row=r, column=1, sticky="ew", pady=6)
            if extra:
                extra.grid(row=r, column=2, padx=(8, 0), pady=6)

        def entry(parent, value, ph=""):
            e = ctk.CTkEntry(parent, height=36, corner_radius=10, font=F(14), placeholder_text=ph)
            if value not in ("", None):
                e.insert(0, str(value))
            return e

        # General
        g = section("General")
        ap = ctk.CTkSegmentedButton(g, values=["System", "Light", "Dark"], font=F(13),
                                    selected_color=ACCENT, selected_hover_color=ACCENT_HOVER,
                                    fg_color=TRACK, unselected_color=TRACK,
                                    unselected_hover_color=SUBTLE_HOVER, text_color=TEXT,
                                    command=lambda v: ctk.set_appearance_mode(v))
        ap.set(S["appearance"])
        row(g, 0, "Appearance", ap)
        self.set_w["appearance"] = ap
        self.set_w["monthly_goal"] = entry(g, S["monthly_goal"])
        row(g, 1, "Monthly goal", self.set_w["monthly_goal"])
        self.set_w["followup_days"] = entry(g, S["followup_days"])
        row(g, 2, "Follow up after (days)", self.set_w["followup_days"])

        # Defaults
        d = section("Defaults for new entries")
        for i, (k, label) in enumerate((("workload", "Workload"), ("method", "Application method"),
                                        ("source", "Source"), ("location", "Location"))):
            self.set_w["def_" + k] = entry(d, S["defaults"].get(k, ""))
            row(d, i, label, self.set_w["def_" + k])
        self.set_w["status_options"] = entry(d, ", ".join(S["status_options"]))
        row(d, 4, "Status options", self.set_w["status_options"])
        self.set_w["rav_options"] = entry(d, ", ".join(S["rav_options"]))
        row(d, 5, "RAV result options", self.set_w["rav_options"])
        ctk.CTkLabel(d, text="Comma separated. The first option is used for new entries.",
                     font=F(12), text_color=MUTED).grid(row=6, column=1, sticky="w")

        # Folders and files
        p = section("Folders and files")
        self.set_w["job_dir"] = entry(p, S["job_dir"], f"automatic: {os.path.dirname(SCRIPT_DIR)}")
        row(p, 0, "Job folder", self.set_w["job_dir"],
            self.button(p, "Browse", lambda: self.browse("job_dir", folder=True), "subtle", width=90, height=36))
        self.set_w["excel_file"] = entry(p, S["excel_file"], "automatic: the only .xlsx in the job folder")
        row(p, 1, "Tracker file", self.set_w["excel_file"],
            self.button(p, "Browse", lambda: self.browse("excel_file", types=[("Excel", "*.xlsx")]),
                        "subtle", width=90, height=36))
        self.set_w["folder_pattern"] = ctk.CTkComboBox(
            p, values=PATTERN_PRESETS, height=36, corner_radius=10, font=F(14),
            button_color=ACCENT, command=lambda v: self.update_preview())
        self.set_w["folder_pattern"].set(S["folder_pattern"])
        self.set_w["folder_pattern"].bind("<KeyRelease>", lambda e: self.update_preview())
        row(p, 2, "Folder name pattern", self.set_w["folder_pattern"])
        self.preview = ctk.CTkLabel(p, text="", font=F(12), text_color=MUTED, justify="left")
        self.preview.grid(row=3, column=1, sticky="w")
        self.update_preview()
        self.tpl_list = list(S["templates"])
        self.tpl_label = ctk.CTkLabel(p, text="", font=F(13), text_color=TEXT, justify="left", anchor="w")
        tbtn = ctk.CTkFrame(p, fg_color="transparent")
        self.button(tbtn, "Add", self.add_templates, "subtle", width=70, height=32).pack(side="left")
        self.button(tbtn, "Clear", self.clear_templates, "subtle", width=70, height=32).pack(side="left", padx=(6, 0))
        row(p, 4, "Template files", self.tpl_label, tbtn)
        self.refresh_templates()

        # Celebration
        c = section("Celebrations")
        sw = ctk.CTkSwitch(c, text="", progress_color=ACCENT)
        if S["sound"]:
            sw.select()
        cs = ctk.CTkSwitch(c, text="Button clicks", progress_color=ACCENT, font=F(13),
                           text_color=TEXT)
        if S.get("click_sounds", True):
            cs.select()
        row(c, 0, "Play sounds", sw)
        self.set_w["sound"] = sw
        self.set_w["click_sounds"] = cs
        for i, (key, label) in enumerate((("sound_file", "Goal sound"),
                                          ("sound_interview", "Interview sound"),
                                          ("sound_offer", "Job offer sound")), start=1):
            self.set_w[key] = entry(c, S[key])
            row(c, i, label, self.set_w[key],
                self.button(c, "Browse", lambda k=key: self.browse(k, types=[("Wave", "*.wav")]),
                            "subtle", width=90, height=36))
        self.test_opts = (["Goal reached"] + [f"+{k} over goal" for k in range(1, len(GOAL_MESSAGES) + 1)]
                          + ["Interview", "Job offer", "All, one after another"])
        tb = ctk.CTkFrame(c, fg_color="transparent")
        self.test_box = ctk.CTkComboBox(tb, values=self.test_opts, height=36, width=230,
                                        corner_radius=10, font=F(14), button_color=ACCENT, state="readonly")
        self.test_box.set("Goal reached")
        self.test_box.pack(side="left")
        self.button(tb, "▶  Test", self.test_celebration, width=100, height=36).pack(side="left", padx=(8, 0))
        row(c, 4, "Test celebration", tb)
        cs.grid(row=0, column=1, sticky="e", pady=6)

        btns = ctk.CTkFrame(self.body, fg_color="transparent")
        btns.pack(fill="x", pady=(14, 0))
        self.button(btns, "Cancel", self.cancel_settings, "subtle", width=120).pack(side="left")
        self.button(btns, "Save settings  ✓", self.save_settings_view, width=180).pack(side="right")

    def update_preview(self):
        pat = self.set_w["folder_pattern"].get()
        old = S["folder_pattern"]
        S["folder_pattern"] = pat
        ex = make_folder_name(today_str(), "Novelis", "Senior Engineer")
        S["folder_pattern"] = old
        self.preview.configure(text=f"Example: {ex}\nTokens: {{yyyy}} {{yy}} {{mm}} {{dd}} {{company}} {{position}}")

    def browse(self, key, folder=False, types=None):
        if folder:
            p = filedialog.askdirectory(parent=self.root)
        else:
            p = filedialog.askopenfilename(parent=self.root, filetypes=types or [])
        if p:
            w = self.set_w[key]
            w.delete(0, "end")
            w.insert(0, os.path.normpath(p))

    def add_templates(self):
        ps = filedialog.askopenfilenames(parent=self.root)
        for p in ps:
            p = os.path.normpath(p)
            if p not in self.tpl_list:
                self.tpl_list.append(p)
        self.refresh_templates()

    def clear_templates(self):
        self.tpl_list = []
        self.refresh_templates()

    def refresh_templates(self):
        txt = "\n".join(os.path.basename(p) for p in self.tpl_list) or "None"
        self.tpl_label.configure(text=txt)

    def test_preset(self, label):
        if label == "Interview":
            p = interview_preset("Novelis")
        elif label == "Job offer":
            p = offer_preset("Novelis")
        elif label == "Goal reached":
            p = goal_preset(0)
        else:
            p = goal_preset(int(label.split()[0].lstrip("+")))
        key = p.get("sound", "sound_file")
        p["sound_path"] = script_rel(self.set_w[key].get().strip()) if self.set_w["sound"].get() else ""
        return p

    def test_celebration(self):
        label = self.test_box.get()
        if label.startswith("All"):
            queue = list(self.test_opts[:-1])

            def nxt():
                if queue:
                    Celebration(self.root, on_done=nxt, **self.test_preset(queue.pop(0)))
            nxt()
        else:
            Celebration(self.root, **self.test_preset(label))

    def cancel_settings(self):
        ctk.set_appearance_mode(S["appearance"])
        self.home()

    def save_settings_view(self):
        w = self.set_w
        try:
            goal = int(w["monthly_goal"].get())
            days = int(w["followup_days"].get())
        except ValueError:
            self.toast("Goal and follow up days must be whole numbers")
            return
        status = [s.strip() for s in w["status_options"].get().split(",") if s.strip()]
        rav = [s.strip() for s in w["rav_options"].get().split(",") if s.strip()]
        if not status or not rav:
            self.toast("Status and RAV options cannot be empty")
            return
        jd = w["job_dir"].get().strip()
        if jd and not os.path.isdir(jd):
            self.toast("Job folder does not exist")
            return
        S.update({
            "appearance": w["appearance"].get(),
            "monthly_goal": goal,
            "followup_days": days,
            "status_options": status,
            "rav_options": rav,
            "job_dir": jd,
            "excel_file": w["excel_file"].get().strip(),
            "folder_pattern": w["folder_pattern"].get().strip() or PATTERN_PRESETS[0],
            "templates": self.tpl_list,
            "sound": bool(w["sound"].get()),
            "click_sounds": bool(w["click_sounds"].get()),
            "sound_file": w["sound_file"].get().strip() or DEFAULT_SETTINGS["sound_file"],
            "sound_interview": w["sound_interview"].get().strip() or DEFAULT_SETTINGS["sound_interview"],
            "sound_offer": w["sound_offer"].get().strip() or DEFAULT_SETTINGS["sound_offer"],
        })
        for k in ("workload", "method", "source", "location"):
            S["defaults"][k] = w["def_" + k].get().strip()
        try:
            save_settings()
        except Exception as e:
            self.error(e)
            return
        ctk.set_appearance_mode(S["appearance"])
        self.home()
        self.toast("Settings saved")


def main():
    ctk.set_appearance_mode(S["appearance"])
    ctk.set_default_color_theme("blue")
    root = ctk.CTk()
    App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
